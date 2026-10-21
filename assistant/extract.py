"""Turn common file types into plain text the assistant can read."""
import html
import io
import re
import zipfile

TEXT_TYPES = (".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".log", ".yaml",
              ".yml", ".ini", ".py", ".js", ".srt")

SUPPORTED = TEXT_TYPES + (".docx", ".pptx", ".xlsx", ".pdf")


def can_read(name: str) -> bool:
    return name.lower().endswith(SUPPORTED)


def _xml_text(xml: str, paragraph_tag: str) -> str:
    xml = xml.replace(f"</{paragraph_tag}>", "\n")
    return html.unescape(re.sub(r"<[^>]+>", "", xml))


def text_from_file(name: str, data: bytes) -> str | None:
    """Plain text of a document, or None if the type isn't supported."""
    ext = "." + name.lower().rsplit(".", 1)[-1] if "." in name else ""
    if ext in TEXT_TYPES:
        return data.decode("utf-8", "replace")
    if ext == ".docx":
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return _xml_text(z.read("word/document.xml").decode("utf-8"), "w:p")
    if ext == ".pptx":
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            slides = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                            key=lambda n: int(re.findall(r"\d+", n)[-1]))
            return "\n\n".join(f"[Slide {i}]\n" + _xml_text(z.read(n).decode("utf-8"), "a:p")
                               for i, n in enumerate(slides, 1))
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"[Sheet {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    out.append(",".join("" if c is None else str(c) for c in row))
        return "\n".join(out)
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    return None
