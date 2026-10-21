"""Writing: reports, essays, letters, CVs... delivered as real Word files."""
import io
import re

from .base import Skill, Tool, schema


def _add_runs(paragraph, text: str) -> None:
    """**bold** and *italic* inside a line."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        elif part:
            paragraph.add_run(part)


def markdown_to_docx(title: str, content: str) -> bytes:
    from docx import Document
    doc = Document()
    doc.add_heading(title, level=0)
    for line in content.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if m := re.match(r"(#{1,4})\s+(.*)", stripped):
            doc.add_heading(m.group(2), level=len(m.group(1)))
        elif m := re.match(r"[-*•]\s+(.*)", stripped):
            _add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif m := re.match(r"\d+[.)]\s+(.*)", stripped):
            _add_runs(doc.add_paragraph(style="List Number"), m.group(1))
        else:
            _add_runs(doc.add_paragraph(), stripped)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def document_create(a, ctx):
    title = a["title"].strip()
    name = re.sub(r"[^\w\- åäöÅÄÖ]+", "", title).strip().replace(" ", "_")[:60] or "document"
    if a.get("format") == "md":
        ctx.files.append((f"{name}.md", f"# {title}\n\n{a['content']}".encode()))
        return f"{name}.md will be sent to the user with your reply."
    ctx.files.append((f"{name}.docx", markdown_to_docx(title, a["content"])))
    return f"{name}.docx (Word) will be sent to the user with your reply."


class WritingSkill(Skill):
    def __init__(self):
        super().__init__(
            name="Writing & documents",
            description="Reports, essays, letters, CVs and more, delivered as Word files.",
            instructions=(
                "For anything longer than a few paragraphs (report, essay, assignment, letter, "
                "CV, cover letter, proposal, summary of a file): if purpose, audience, length or "
                "language is unclear, ask ONE short question; for long pieces agree on an outline "
                "first. Write in a clear, natural voice, with real sources from web_search where "
                "facts matter (cite them; use APA for academic work unless told otherwise). "
                "Deliver finished documents with document_create (Word by default) and keep the "
                "chat reply to a short summary. Also: proofread and give concrete feedback on the "
                "user's own texts, rewrite in other tones, shorten, or translate them."
            ),
            tools=[
                Tool("document_create",
                     "Create a Word (or Markdown) file and send it to the user in the chat. "
                     "Content uses simple markdown: # headings, - bullets, 1. lists, **bold**.",
                     schema({"title": {"type": "string"},
                             "content": {"type": "string", "description": "The full text, markdown"},
                             "format": {"type": "string", "enum": ["docx", "md"]}},
                            ["title", "content"]), document_create),
            ],
        )
