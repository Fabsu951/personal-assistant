"""Offline tests: a fake Claude drives the agent loop, no API keys needed."""
import asyncio
from datetime import datetime, timedelta
from types import SimpleNamespace as NS

from assistant.agent import Agent
from assistant.config import Config
from assistant.db import DB
from assistant.skills import load_skills
from assistant.skills.tasks import due_reminders


def text(t):
    return NS(type="text", text=t)


def tool_use(id, name, inp):
    return NS(type="tool_use", id=id, name=name, input=inp)


class FakeClient:
    def __init__(self, script):
        self.script = list(script)
        self.calls = []
        self.messages = self

    async def create(self, **kw):
        self.calls.append(kw)
        content, stop = self.script.pop(0)
        return NS(content=content, stop_reason=stop)


def make(tmp_path, script):
    cfg = Config()
    cfg.data_dir = tmp_path
    db = DB(tmp_path / "t.db")
    skills = load_skills(cfg.enabled_skills, db, cfg)
    client = FakeClient(script)
    return Agent(cfg, db, skills, client=client), db, client


def test_tool_runs_and_history_saved(tmp_path):
    agent, db, client = make(tmp_path, [
        ([tool_use("t1", "task_add", {"title": "Renew passport"})], "tool_use"),
        ([text("Added it!")], "end_turn"),
    ])
    reply = asyncio.run(agent.handle(1, "add renew passport"))
    assert reply.text == "Added it!"
    assert db.query("SELECT title FROM tasks")[0]["title"] == "Renew passport"
    tool_result = client.calls[1]["messages"][-1]["content"][0]
    assert "Task #1 added" in tool_result["content"]
    assert [m["role"] for m in db.history(1, 10)] == ["user", "assistant"]
    # tools sent to the API include web search + all client tools
    names = [t["name"] for t in client.calls[0]["tools"]]
    assert "web_search" in names and "gmail_send" in names


def test_approval_flow(tmp_path):
    email = {"to": "a@b.com", "subject": "Hi", "body": "Hello"}
    agent, db, _ = make(tmp_path, [
        ([tool_use("t1", "gmail_send", email)], "tool_use"),
        ([text("Drafted, waiting for approval.")], "end_turn"),
    ])
    reply = asyncio.run(agent.handle(1, "email a@b.com"))
    assert reply.pending_ids == [1]
    row = db.query("SELECT * FROM pending_actions")[0]
    assert row["status"] == "pending" and "a@b.com" in row["summary"]
    # Reject -> nothing executed
    assert "Rejected" in asyncio.run(agent.resolve_action(1, approve=False))
    assert "Already" in asyncio.run(agent.resolve_action(1, approve=True))


def test_approve_runs_tool(tmp_path):
    agent, db, _ = make(tmp_path, [])
    sent = {}
    agent.tools["gmail_send"].handler = lambda a, c: sent.update(a) or "Email sent."
    pid = db.execute("INSERT INTO pending_actions(chat_id, tool, args, summary) VALUES(1,'gmail_send','{\"to\":\"x\"}','s')")
    out = asyncio.run(agent.resolve_action(pid, approve=True))
    assert out.startswith("Done") and sent == {"to": "x"}


def test_tool_error_is_reported_not_crash(tmp_path):
    agent, db, client = make(tmp_path, [
        ([tool_use("t1", "gmail_search", {"query": "x"})], "tool_use"),
        ([text("Google isn't connected yet.")], "end_turn"),
    ])
    reply = asyncio.run(agent.handle(1, "emails?"))
    assert "connected" in reply.text
    assert "ERROR" in client.calls[1]["messages"][-1]["content"][0]["content"]


def test_reminders_due(tmp_path):
    agent, db, client = make(tmp_path, [])
    past = (datetime.now().astimezone() - timedelta(minutes=1)).isoformat()
    db.execute("INSERT INTO reminders(chat_id, text, due_at) VALUES(1,'call mom',?)", (past,))
    assert [r["text"] for r in due_reminders(db)] == ["call mom"]
    assert due_reminders(db) == []


def test_history_alternates(tmp_path):
    agent, db, _ = make(tmp_path, [])
    db.add_message(1, "assistant", "(approved)")
    db.add_message(1, "user", "a")
    db.add_message(1, "assistant", "b")
    db.add_message(1, "assistant", "(action done)")
    msgs = agent._history_messages(1)
    assert [m["role"] for m in msgs] == ["user", "assistant"]


def test_expenses(tmp_path):
    agent, db, _ = make(tmp_path, [])
    from assistant.skills.base import ToolContext
    ctx = ToolContext(db=db, config=agent.config, chat_id=1)
    agent.tools["expense_add"].handler({"amount": 250, "category": "food"}, ctx)
    s = agent.tools["expense_summary"].handler({}, ctx)
    assert s["totals_by_category"][0]["total"] == 250


def test_prompt_caching(tmp_path):
    agent, db, client = make(tmp_path, [([text("Hi")], "end_turn")])
    asyncio.run(agent.handle(1, "hello"))
    call = client.calls[0]
    assert call["cache_control"] == {"type": "ephemeral"}
    stable, now = call["system"]
    assert stable["cache_control"] == {"type": "ephemeral", "ttl": "1h"}
    # the time changes every minute, so it must not be inside the cached block
    assert "Current date and time" not in stable["text"] and "Current date and time" in now["text"]
    assert "cache_control" not in now


def test_system_prompt_identical_within_turn(tmp_path):
    agent, db, client = make(tmp_path, [
        ([tool_use("t1", "task_list", {})], "tool_use"),
        ([text("Nothing to do.")], "end_turn"),
    ])
    asyncio.run(agent.handle(1, "tasks?"))
    assert client.calls[0]["system"] == client.calls[1]["system"]


class FakeCalendar:
    """Just enough of the Google Calendar client for calendar_list_events."""
    def __init__(self, calendars, events):
        self.calendars, self.events_by_cal, self.inserted = calendars, events, None

    def calendarList(self):
        return NS(list=lambda pageToken=None: NS(execute=lambda: {"items": self.calendars}))

    def events(self):
        def list_(calendarId, **kw):
            if calendarId not in self.events_by_cal:
                raise RuntimeError("forbidden")
            return NS(execute=lambda: {"items": self.events_by_cal[calendarId]})

        def insert(calendarId, body, sendUpdates):
            self.inserted = calendarId
            return NS(execute=lambda: {"htmlLink": "link"})
        return NS(list=list_, insert=insert)


def test_calendar_reads_all_calendars(tmp_path, monkeypatch):
    from assistant.skills import google_workspace as gw
    from assistant.skills.base import ToolContext
    agent, db, _ = make(tmp_path, [])
    cal = FakeCalendar(
        [{"id": "me@x.com", "summary": "me@x.com", "primary": True},
         {"id": "fam123", "summary": "Familj"},
         {"id": "broken", "summary": "Broken"}],
        {"me@x.com": [{"id": "a", "summary": "Gym",
                       "start": {"dateTime": "2026-10-21T18:00:00+08:00"},
                       "end": {"dateTime": "2026-10-21T19:00:00+08:00"}}],
         "fam123": [{"id": "b", "summary": "Dinner with mom",
                     "start": {"dateTime": "2026-10-21T10:00:00Z"},
                     "end": {"dateTime": "2026-10-21T11:00:00Z"}},
                    {"id": "c", "summary": "Birthday", "start": {"date": "2026-10-21"},
                     "end": {"date": "2026-10-22"}}]})
    monkeypatch.setattr(gw, "_svc", lambda ctx, name, version: cal)
    ctx = ToolContext(db=db, config=agent.config, chat_id=1)
    out = gw.calendar_list_events({"start": "2026-10-21", "end": "2026-10-22"}, ctx)
    assert [e.get("title") for e in out[:3]] == ["Birthday", "Gym", "Dinner with mom"]
    assert out[2]["calendar"] == "Familj" and out[2]["calendar_id"] == "fam123"
    assert "Broken" in out[3]["note"]
    gw.calendar_create_event({"title": "Picnic", "start": "2026-10-25T12:00", "calendar": "familj"}, ctx)
    assert cal.inserted == "fam123"


def test_fitness(tmp_path):
    from assistant.skills.base import ToolContext
    agent, db, _ = make(tmp_path, [])
    ctx = ToolContext(db=db, config=agent.config, chat_id=1)
    t = agent.tools
    t["program_set"].handler({"program": "Mon: legs. Thu: upper body."}, ctx)
    assert "legs" in t["program_get"].handler({}, ctx)["program"]
    t["workout_log"].handler({"kind": "legs", "date": "2000-01-01", "exercises": [
        {"exercise": "Squat", "sets": 3, "reps": 5, "weight_kg": 100}]}, ctx)
    t["workout_log"].handler({"kind": "legs", "exercises": [
        {"exercise": "squat", "sets": 3, "reps": 5, "weight_kg": 102.5}]}, ctx)
    assert [r["weight_kg"] for r in t["exercise_progress"].handler({"exercise": "squat"}, ctx)] == [100, 102.5]
    recent = t["workout_history"].handler({"days": 7}, ctx)
    assert len(recent) == 1 and recent[0]["exercises"][0]["weight_kg"] == 102.5


def test_voice_needs_key():
    import pytest
    from assistant import voice
    cfg = Config()
    cfg.openai_api_key = ""
    with pytest.raises(RuntimeError, match="OPENAI_API_KEY"):
        asyncio.run(voice.transcribe(cfg, b"audio"))


def test_voice_transcribes(monkeypatch):
    import httpx
    from assistant import voice
    seen = {}

    def handler(request):
        seen["auth"] = request.headers["authorization"]
        return httpx.Response(200, json={"text": " hej Ragnar "})

    real = httpx.AsyncClient
    monkeypatch.setattr(voice.httpx, "AsyncClient",
                        lambda **kw: real(transport=httpx.MockTransport(handler), **kw))
    cfg = Config()
    cfg.openai_api_key = "sk-test"
    assert asyncio.run(voice.transcribe(cfg, b"audio")) == "hej Ragnar"
    assert seen["auth"] == "Bearer sk-test"


def test_blank_settings_use_defaults(monkeypatch):
    for name in ["MODEL", "TIMEZONE", "HISTORY_TURNS", "DATA_DIR", "ENABLED_SKILLS", "OWNER_TELEGRAM_ID"]:
        monkeypatch.setenv(name, "")
    cfg = Config()
    assert cfg.model and cfg.timezone == "Asia/Taipei" and cfg.history_turns == 20
    assert "google" in cfg.enabled_skills and cfg.owner_id == 0 and str(cfg.data_dir) == "data"


def test_body_status(tmp_path):
    from assistant.skills.base import ToolContext
    agent, db, _ = make(tmp_path, [])
    ctx = ToolContext(db=db, config=agent.config, chat_id=1)
    agent.tools["status_log"].handler({"sleep_hours": 6, "injuries": "left knee aches"}, ctx)
    agent.tools["status_log"].handler({"weight_kg": 80.5, "energy": 8}, ctx)
    rows = agent.tools["status_history"].handler({"days": 3}, ctx)
    assert rows[0]["injuries"] == "left knee aches" and "weight_kg" not in rows[0]
    assert rows[1]["weight_kg"] == 80.5


def test_strava(tmp_path, monkeypatch):
    import httpx
    from assistant import strava
    agent, db, _ = make(tmp_path, [])
    cfg = agent.config
    cfg.strava_client_id, cfg.strava_client_secret = "123", "secret"
    assert "state=strava" in strava.start(cfg)
    assert strava.is_strava_redirect("http://localhost/?state=strava&code=abc&scope=read")
    assert not strava.is_strava_redirect("http://localhost/?code=abc&scope=https://www.googleapis.com/x")

    calls = []

    def fake_post(url, data, timeout):
        calls.append(data["grant_type"])
        return httpx.Response(200, json={"access_token": f"tok{len(calls)}", "refresh_token": "r",
                                         "expires_at": 0 if len(calls) == 1 else 9999999999})

    def fake_get(url, headers, params, timeout):
        assert headers["Authorization"] == "Bearer tok2"  # expired token was refreshed first
        return httpx.Response(200, json=[{"name": "Morning run", "sport_type": "Run",
                                          "start_date_local": "2026-10-19T07:00:00Z",
                                          "moving_time": 1800, "distance": 6000,
                                          "average_heartrate": 150}])

    monkeypatch.setattr(strava.httpx, "post", fake_post)
    monkeypatch.setattr(strava.httpx, "get", fake_get)
    strava.finish(db, cfg, "http://localhost/?state=strava&code=abc")
    run = strava.activities(db, cfg, 7)[0]
    assert calls == ["authorization_code", "refresh_token"]
    assert run["distance_km"] == 6.0 and run["pace_min_per_km"] == "5:00" and run["moving_min"] == 30
    assert db.get("strava_token")["access_token"] == "tok2"


def test_strava_not_connected(tmp_path):
    agent, db, client = make(tmp_path, [
        ([tool_use("t1", "strava_activities", {})], "tool_use"),
        ([text("Strava isn't connected.")], "end_turn"),
    ])
    asyncio.run(agent.handle(1, "runs?"))
    assert "connect_strava" in client.calls[1]["messages"][-1]["content"][0]["content"]


def test_voice_modes():
    from assistant.voice import for_speech, wants_voice
    assert wants_voice("auto", spoken=True) and not wants_voice("auto", spoken=False)
    assert wants_voice("always", spoken=False) and not wants_voice("off", spoken=True)
    assert for_speech("**Klart!** Se https://x.se/a?b=1 nu") == "Klart! Se  nu"


def test_voice_speaks(monkeypatch):
    import json
    import httpx
    from assistant import voice
    seen = {}

    def handler(request):
        seen.update(json.loads(request.content))
        return httpx.Response(200, content=b"OggS-audio")

    real = httpx.AsyncClient
    monkeypatch.setattr(voice.httpx, "AsyncClient",
                        lambda **kw: real(transport=httpx.MockTransport(handler), **kw))
    cfg = Config()
    cfg.openai_api_key = "sk-test"
    assert asyncio.run(voice.speak(cfg, "Hej *Fabian*")) == b"OggS-audio"
    assert seen["input"] == "Hej Fabian" and seen["response_format"] == "opus"


def test_voice_settings_tool(tmp_path, monkeypatch):
    import json
    import httpx
    from assistant import voice
    from assistant.skills.base import ToolContext
    agent, db, _ = make(tmp_path, [])
    ctx = ToolContext(db=db, config=agent.config, chat_id=1)
    out = agent.tools["voice_settings"].handler({"voice": "nova", "style": "Very happy"}, ctx)
    assert out["voice"] == "nova" and out["style"] == "Very happy" and out["replies"] == "auto"
    seen = {}

    def handler(request):
        seen.update(json.loads(request.content))
        return httpx.Response(200, content=b"ogg")

    real = httpx.AsyncClient
    monkeypatch.setattr(voice.httpx, "AsyncClient",
                        lambda **kw: real(transport=httpx.MockTransport(handler), **kw))
    agent.config.openai_api_key = "sk-test"
    asyncio.run(voice.speak(agent.config, "Hej", voice=db.get("tts_voice"), style=db.get("tts_style")))
    assert seen["voice"] == "nova" and seen["instructions"].startswith("Very happy")


def test_voice_skill_always_loaded(tmp_path):
    cfg = Config()
    db = DB(tmp_path / "t.db")
    names = [s.name for s in load_skills(["tasks"], db, cfg)]
    assert names == ["Tasks & reminders", "Voice"]


def test_spoken_turn_is_fast(tmp_path):
    agent, db, client = make(tmp_path, [([text("Hej!")], "end_turn"), ([text("Hej igen")], "end_turn")])
    asyncio.run(agent.handle(1, "hej", spoken=True))
    fast = client.calls[0]
    assert fast["output_config"] == {"effort": "low"} and "read aloud" in fast["system"][1]["text"]
    asyncio.run(agent.handle(1, "hej"))
    normal = client.calls[1]
    assert "output_config" not in normal and "read aloud" not in normal["system"][1]["text"]
    assert fast["system"][0] == normal["system"][0]  # cached part unchanged


def _signed_init_data(token, user_id, auth_date=None):
    import hashlib
    import hmac
    import json
    import time
    from urllib.parse import urlencode
    fields = {"auth_date": str(auth_date or int(time.time())), "query_id": "q1",
              "user": json.dumps({"id": user_id, "first_name": "Fabian"})}
    check = "\n".join(f"{k}={v}" for k, v in sorted(fields.items()))
    secret = hmac.new(b"WebAppData", token.encode(), hashlib.sha256).digest()
    fields["hash"] = hmac.new(secret, check.encode(), hashlib.sha256).hexdigest()
    return urlencode(fields)


def test_web_auth():
    from assistant import web
    good = _signed_init_data("123:abc", 42)
    assert web.check_init_data(good, "123:abc") == 42
    assert web.check_init_data(good, "999:other") is None                  # wrong bot
    assert web.check_init_data(good.replace("42", "43"), "123:abc") is None  # tampered
    assert web.check_init_data(_signed_init_data("123:abc", 42, auth_date=1), "123:abc") is None  # old
    key = web.make_key("123:abc", 42)
    assert web.check_key(key, "123:abc", 42) and not web.check_key(key, "123:abc", 43)
    assert not web.check_key("1.deadbeef", "123:abc", 42)


def test_web_talk(tmp_path, monkeypatch):
    from aiohttp.test_utils import TestClient, TestServer
    from assistant import voice, web

    agent, db, client = make(tmp_path, [([text("Du har gym kl 18.")], "end_turn")])
    agent.config.telegram_token, agent.config.owner_id = "123:abc", 42

    async def fake_transcribe(cfg, audio, filename):
        assert audio == b"sound" and filename == "talk.webm"
        return "Vad har jag idag?"

    async def fake_speak(cfg, text, voice=None, style=None, fmt="opus"):
        assert fmt == "mp3"
        return b"mp3-bytes"

    monkeypatch.setattr(voice, "transcribe", fake_transcribe)
    monkeypatch.setattr(voice, "speak", fake_speak)
    bot = NS(config=agent.config, agent=agent, db=db)

    async def run():
        async with TestClient(TestServer(web.WebApp(bot).app)) as c:
            page = await c.get("/talk")
            assert page.status == 200 and "Prata med Ragnar" in await page.text()
            denied = await c.post("/api/talk", data=b"sound", headers={"Content-Type": "audio/webm"})
            assert denied.status == 403
            ok = await c.post("/api/talk", data=b"sound", headers={
                "Content-Type": "audio/webm", "X-Init-Data": _signed_init_data("123:abc", 42)})
            return ok.status, await ok.json()

    status, data = asyncio.run(run())
    assert status == 200 and data["heard"] == "Vad har jag idag?"
    assert data["reply"] == "Du har gym kl 18." and data["audio"] == "bXAzLWJ5dGVz"
    assert "read aloud" in client.calls[0]["system"][1]["text"]  # short spoken answer
    assert [m["role"] for m in db.history(42, 10)] == ["user", "assistant"]


def test_extract_office_files(tmp_path):
    import io
    import zipfile
    from openpyxl import Workbook
    from assistant.extract import can_read, text_from_file

    def office(files):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            for name, xml in files.items():
                z.writestr(name, xml)
        return buf.getvalue()

    docx = office({"word/document.xml": "<w:document><w:p><w:r><w:t>Hej &amp; välkommen</w:t></w:r></w:p>"
                                        "<w:p><w:t>Rad två</w:t></w:p></w:document>"})
    assert text_from_file("Brev.DOCX", docx).splitlines() == ["Hej & välkommen", "Rad två"]
    pptx = office({"ppt/slides/slide2.xml": "<a:p><a:t>Andra</a:t></a:p>",
                   "ppt/slides/slide1.xml": "<a:p><a:t>Första</a:t></a:p>"})
    assert text_from_file("deck.pptx", pptx).index("Första") < text_from_file("deck.pptx", pptx).index("Andra")
    wb = Workbook()
    wb.active.append(["Datum", "Belopp"])
    wb.active.append(["2026-10-01", 250])
    buf = io.BytesIO()
    wb.save(buf)
    assert "2026-10-01,250" in text_from_file("budget.xlsx", buf.getvalue())
    assert text_from_file("video.mp4", b"x") is None and not can_read("video.mp4") and can_read("a.PDF")


def test_flights_search(tmp_path, monkeypatch):
    import httpx
    from assistant.skills import travel
    from assistant.skills.base import ToolContext
    agent, db, _ = make(tmp_path, [])
    ctx = ToolContext(db=db, config=agent.config, chat_id=1)
    assert "SERPAPI_API_KEY" in str(_raises(lambda: travel.flights_search(
        {"from": "TPE", "to": "ARN", "depart_date": "2026-12-10"}, ctx)))
    agent.config.serpapi_api_key = "k"
    seen = {}

    def leg(a, b, t1, t2, no, mins):
        return {"departure_airport": {"id": a, "time": f"2026-12-10 {t1}"},
                "arrival_airport": {"id": b, "time": f"2026-12-10 {t2}"},
                "airline": no[:2], "flight_number": no, "duration": mins}

    def fake_get(url, params, timeout):
        seen.update(params)
        return httpx.Response(200, json={
            "best_flights": [{"price": 14000, "total_duration": 900, "flights": [leg("TPE", "ARN", "08:00", "18:00", "BR 67", 900)]}],
            "other_flights": [{"price": 9800, "total_duration": 1300,
                               "flights": [leg("TPE", "IST", "23:00", "06:00", "TK 25", 780),
                                           leg("IST", "ARN", "09:00", "11:30", "TK 1795", 210)],
                               "layovers": [{"id": "IST", "duration": 180}]}],
            "price_insights": {"price_level": "low", "typical_price_range": [11000, 16000]},
            "search_metadata": {"google_flights_url": "https://www.google.com/travel/flights?x"}})

    monkeypatch.setattr(travel.httpx, "get", fake_get)
    out = travel.flights_search({"from": "tpe, tsa", "to": "ARN", "depart_date": "2026-12-10",
                                 "return_date": "2027-01-05", "max_stops": "max_1"}, ctx)
    assert seen["departure_id"] == "TPE,TSA" and seen["type"] == 1 and seen["stops"] == 2
    assert seen["currency"] == "SEK" and seen["sort_by"] == 2
    cheapest = out["cheapest"][0]
    assert cheapest["price"] == 9800 and cheapest["stops"] == 1 and cheapest["layovers"] == ["IST 3h00"]
    assert cheapest["legs"][0] == "TPE 23:00 → IST 06:00 (TK 25, 13h00)"
    assert out["price_level"] == "low" and out["google_flights_link"].startswith("https://")


def _raises(fn):
    try:
        fn()
    except Exception as e:
        return e
    raise AssertionError("expected an error")


def test_onedrive(tmp_path, monkeypatch):
    import io
    import time
    import zipfile
    import httpx
    from assistant import microsoft
    from assistant.skills.base import ToolContext
    agent, db, _ = make(tmp_path, [])
    cfg = agent.config
    cfg.ms_client_id = "app-id"
    ctx = ToolContext(db=db, config=cfg, chat_id=1)
    assert "connect_onedrive" in str(_raises(lambda: agent.tools["onedrive_search"].handler({"query": "x"}, ctx)))

    # device-code login: pending once, then success
    answers = [httpx.Response(400, json={"error": "authorization_pending"}),
               httpx.Response(200, json={"access_token": "A1", "refresh_token": "R1", "expires_in": 3600})]

    class FakeAsyncClient:
        def __init__(self, **kw): pass
        async def __aenter__(self): return self
        async def __aexit__(self, *a): pass
        async def post(self, url, data):
            assert data["device_code"] == "dev" and url.endswith("/token")
            return answers.pop(0)

    async def no_sleep(s): pass
    monkeypatch.setattr(microsoft.httpx, "AsyncClient", FakeAsyncClient)
    monkeypatch.setattr(microsoft.asyncio, "sleep", no_sleep)
    asyncio.run(microsoft.wait_for_login(db, cfg, {"device_code": "dev", "interval": 1, "expires_in": 60}))
    assert db.get(microsoft.TOKEN_KEY)["access_token"] == "A1"

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", "<w:p><w:t>Hyresavtal</w:t></w:p>")
    item = {"id": "1", "name": "Avtal.docx", "size": 100, "webUrl": "https://onedrive/x",
            "file": {}, "parentReference": {"path": "/drive/root:/Dokument"}}

    def fake_get(url, params=None, follow_redirects=True, timeout=60, headers=None):
        assert headers["Authorization"] == "Bearer A1"
        if "search(q='hyres''avtal')" in url:
            return httpx.Response(200, json={"value": [item]})
        if url.endswith("/items/1/content"):
            return httpx.Response(200, content=buf.getvalue())
        if url.endswith("/items/1"):
            return httpx.Response(200, json=item)
        raise AssertionError(url)

    monkeypatch.setattr(microsoft.httpx, "get", fake_get)
    found = agent.tools["onedrive_search"].handler({"query": "hyres'avtal"}, ctx)
    assert found[0]["name"] == "Avtal.docx" and found[0]["folder"] == "/Dokument"
    read = agent.tools["onedrive_read"].handler({"file_id": "1"}, ctx)
    assert read["content"].strip() == "Hyresavtal"
    db.set(microsoft.TOKEN_KEY, {**db.get(microsoft.TOKEN_KEY), "expires_at": time.time() - 1})
    monkeypatch.setattr(microsoft.httpx, "post", lambda url, timeout, data: httpx.Response(
        200, json={"access_token": "A1", "refresh_token": "R2", "expires_in": 3600}))
    agent.tools["onedrive_search"].handler({"query": "hyres'avtal"}, ctx)  # refreshed first
    assert db.get(microsoft.TOKEN_KEY)["refresh_token"] == "R2"


def test_photo_attachment(tmp_path):
    agent, db, client = make(tmp_path, [([text("Det står 'Nötkött nudlar'.")], "end_turn")])
    db.add_message(1, "user", "hej")  # an unanswered earlier message gets merged in
    img = {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": "AAA="}}
    asyncio.run(agent.handle(1, "Översätt", attachments=[img]))
    sent = client.calls[0]["messages"][-1]["content"]
    assert sent[0] == img and sent[1] == {"type": "text", "text": "hej\n\nÖversätt"}
    assert db.history(1, 10)[-2]["content"] == "Översätt [attached 1 file(s)]"


def _ctx(tmp_path):
    from assistant.skills.base import ToolContext
    agent, db, _ = make(tmp_path, [])
    return agent, ToolContext(db=db, config=agent.config, chat_id=1)


def test_tool_names_unique_and_valid(tmp_path):
    import re
    from assistant.skills import AVAILABLE
    cfg = Config()
    names = [t.name for s in load_skills(list(AVAILABLE), DB(tmp_path / "t.db"), cfg) for t in s.tools]
    assert len(names) == len(set(names))
    assert all(re.fullmatch(r"[a-zA-Z0-9_-]{1,64}", n) for n in names)


def test_document_is_sent_as_file(tmp_path):
    import io
    from docx import Document
    agent, db, client = make(tmp_path, [
        ([tool_use("t1", "document_create", {"title": "Rapport: Bioinspirerad design",
                                             "content": "# Inledning\nText med **fet** stil.\n- punkt ett\n1. steg ett"})], "tool_use"),
        ([text("Här är rapporten.")], "end_turn"),
    ])
    reply = asyncio.run(agent.handle(1, "skriv en rapport"))
    (name, data), = reply.files
    assert name == "Rapport_Bioinspirerad_design.docx"
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    assert texts[:3] == ["Rapport: Bioinspirerad design", "Inledning", "Text med fet stil."]
    assert doc.paragraphs[2].runs[1].bold and "punkt ett" in texts and "steg ett" in texts


def test_flashcards_spaced_repetition(tmp_path):
    from assistant.skills.study import next_review
    agent, ctx = _ctx(tmp_path)
    t = agent.tools
    t["flashcards_add"].handler({"deck": "Mandarin", "cards": [{"front": "謝謝", "back": "xièxie – tack"},
                                                               {"front": "你好", "back": "nǐ hǎo – hej"}]}, ctx)
    due = t["flashcards_due"].handler({"deck": "mandarin"}, ctx)
    assert len(due) == 2
    t["flashcard_grade"].handler({"card_id": due[0]["id"], "grade": "good"}, ctx)
    assert len(t["flashcards_due"].handler({}, ctx)) == 1
    assert t["flashcard_decks"].handler({}, ctx)[0] == {"deck": "Mandarin", "cards": 2, "due_now": 1}
    assert next_review(0, 2.5, 0, "good")[0] == 1 and next_review(1, 2.5, 1, "good")[0] == 3
    assert next_review(3, 2.5, 2, "good")[0] == 8 and next_review(8, 2.5, 3, "again")[:1] == (0,)


def test_ideas_content_career(tmp_path):
    agent, ctx = _ctx(tmp_path)
    t = agent.tools
    t["idea_save"].handler({"title": "App som slår ihop 3 appar", "tags": "app,video"}, ctx)
    t["idea_update"].handler({"idea_id": 1, "add_details": "Börja med kalender", "status": "exploring"}, ctx)
    idea = t["idea_list"].handler({"query": "video"}, ctx)[0]
    assert idea["status"] == "exploring" and "kalender" in idea["details"]
    t["content_add"].handler({"title": "3 appar → 1", "platform": "TikTok", "publish_date": "2026-10-24"}, ctx)
    t["content_update"].handler({"content_id": 1, "status": "filmed", "script": "Hook: ..."}, ctx)
    assert t["content_list"].handler({}, ctx)[0]["status"] == "filmed"
    assert t["content_script"].handler({"content_id": 1}, ctx)["script"] == "Hook: ..."
    t["content_update"].handler({"content_id": 1, "status": "published"}, ctx)
    assert t["content_list"].handler({}, ctx) == "Nothing in the content plan."
    t["application_add"].handler({"company": "TSMC", "role": "Intern", "deadline": "2026-11-15"}, ctx)
    t["application_update"].handler({"application_id": 1, "status": "applied", "add_note": "Skickat CV"}, ctx)
    app = t["application_list"].handler({}, ctx)[0]
    assert app["status"] == "applied" and app["applied_on"] and "Skickat CV" in app["notes"]


def test_budget_and_savings(tmp_path):
    from assistant.skills.base import today
    agent, ctx = _ctx(tmp_path)
    t = agent.tools
    t["budget_set"].handler({"category": "Food", "monthly_limit": 6000, "currency": "TWD"}, ctx)
    t["expense_add"].handler({"amount": 4800, "category": "food", "currency": "TWD",
                              "date": today(ctx).isoformat()}, ctx)
    b = t["budget_status"].handler({}, ctx)["budgets"][0]
    assert b["spent"] == 4800 and b["used_percent"] == 80 and b["left"] == 1200
    t["savings_goal_add"].handler({"name": "Japan", "target": 30000, "deadline": "2099-01-01"}, ctx)
    t["savings_goal_update"].handler({"goal_id": 1, "add": 3000}, ctx)
    g = t["savings_goals"].handler({}, ctx)[0]
    assert g["saved"] == 3000 and g["percent"] == 10 and g["needed_per_month"] > 0


def test_nutrition_and_shopping(tmp_path):
    agent, ctx = _ctx(tmp_path)
    t = agent.tools
    t["meal_log"].handler({"description": "牛肉麵", "kcal": 650, "protein_g": 35}, ctx)
    t["meal_log"].handler({"description": "Protein shake", "kcal": 150, "protein_g": 25}, ctx)
    day = t["nutrition_summary"].handler({}, ctx)["per_day"][0]
    assert day["kcal"] == 800 and day["protein_g"] == 60 and day["meals"] == 2
    t["shopping_add"].handler({"items": ["ägg", "kyckling", "ris"]}, ctx)
    t["shopping_done"].handler({"item_ids": [1]}, ctx)
    assert [i["item"] for i in t["shopping_list"].handler({}, ctx)] == ["kyckling", "ris"]
    t["shopping_done"].handler({"all": True}, ctx)
    assert t["shopping_list"].handler({}, ctx) == "Shopping list is empty."


def test_habits_and_goals(tmp_path):
    from datetime import timedelta
    from assistant.skills.base import today
    from assistant.skills.habits import streak
    agent, ctx = _ctx(tmp_path)
    t = agent.tools
    t["habit_add"].handler({"name": "Läsa 20 min"}, ctx)
    now = today(ctx)
    for d in (1, 2):
        t["habit_done"].handler({"habits": ["läsa"], "date": (now - timedelta(days=d)).isoformat()}, ctx)
    h = t["habit_status"].handler({}, ctx)[0]
    assert h["streak_days"] == 2 and not h["done_today"]
    assert t["habit_done"].handler({"habits": ["Läsa", "meditera"]}, ctx) == {"checked": ["Läsa 20 min"], "not_found": ["meditera"]}
    assert t["habit_status"].handler({}, ctx)[0]["streak_days"] == 3
    assert streak(set(), now) == 0
    t["goal_add"].handler({"title": "120 kg knäböj", "target_date": "2027-06-01"}, ctx)
    t["goal_update"].handler({"goal_id": 1, "progress": 40, "note": "105 kg idag"}, ctx)
    g = t["goal_list"].handler({}, ctx)[0]
    assert g["progress"] == 40 and "105 kg" in g["recent_log"]


def test_people_and_birthdays(tmp_path):
    from datetime import date
    from assistant.skills.base import today
    from assistant.skills.people import days_until_birthday
    agent, ctx = _ctx(tmp_path)
    t = agent.tools
    now = today(ctx)
    soon = date(2000, now.month, now.day) if (now.month, now.day) != (2, 29) else date(2000, 3, 1)
    t["contact_save"].handler({"name": "Mamma", "relation": "mother", "birthday": soon.isoformat(),
                               "add_gift_idea": "Teservis"}, ctx)
    t["contact_save"].handler({"name": "mamma", "add_note": "Gillar trädgård"}, ctx)  # same person
    c = t["contact_find"].handler({"query": "trädgård"}, ctx)[0]
    assert c["relation"] == "mother" and "Teservis" in c["gift_ideas"]
    b = t["birthdays_upcoming"].handler({"days": 7}, ctx)[0]
    assert b["in_days"] == 0 and b["turns"] == now.year - 2000
    assert t["contacts_neglected"].handler({}, ctx)[0]["name"] == "Mamma"
    t["contact_touch"].handler({"name": "Mamma", "note": "Pratade om jul"}, ctx)
    assert t["contacts_neglected"].handler({}, ctx) == "You're in touch with everyone."
    assert days_until_birthday("02-29", date(2027, 2, 28)) == 1 and days_until_birthday("12-31", date(2026, 1, 1)) == 364


def test_journal(tmp_path):
    agent, ctx = _ctx(tmp_path)
    t = agent.tools
    t["journal_add"].handler({"entry": "Stressig dag inför tentan", "mood": 4, "tags": "plugg"}, ctx)
    t["journal_add"].handler({"entry": "Bra träning", "mood": 8}, ctx)
    assert len(t["journal_read"].handler({}, ctx)) == 2
    assert t["journal_read"].handler({"query": "tentan"}, ctx)[0]["mood"] == 4
    trend = t["mood_trend"].handler({}, ctx)
    assert trend["average"] == 6.0 and trend["lowest"] == 6.0  # one day, averaged
